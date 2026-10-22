"""Fill the "Template Gambaran sistem.docx" with this thesis' system overview.

    python docs/gambaran-sistem/build_docx.py <template.docx>

Writes Gambaran_Sistem_UWB_Positioning.docx next to this script, using the
PNGs produced by gen_diagrams.py.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "Gambaran_Sistem_UWB_Positioning.docx"

JUDUL = ("Implementasi Indoor Positioning Berbasis Ultra-Wideband Menggunakan Extended "
         "Kalman Filter pada Smart Trolley di Dalam Smart Market")
REFERENSI = [
    "[1]\tD. Arbula dan S. Ljubic, \"Indoor Localization Based on Infrared Angle of Arrival "
    "Sensor Network,\" Sensors, vol. 20, no. 21, 6278, 2020, doi: 10.3390/s20216278.",
    "[2]\tS. Vandermeeren dan H. Steendam, \"PDR/UWB Based Positioning of a Shopping Cart,\" "
    "IEEE Sensors Journal, vol. 21, no. 9, hlm. 10864–10878, 2021, "
    "doi: 10.1109/JSEN.2021.3060110.",
    "[3]\tH.-S. Duan dkk., \"Intelligent Supermarket Shopping Guide System Based on UWB,\" "
    "International Journal of Frontiers in Engineering Technology, vol. 4, no. 6, 2022, "
    "doi: 10.25236/IJFET.2022.040605.",
    "[4]\tX. Liang, Y. Qi, S. Xi, G. Meng, dan Z. Wang, \"Design and development of full "
    "self-service smart shopping system for unmanned supermarket,\" Journal of King Saud "
    "University – Computer and Information Sciences, vol. 37, no. 6, 152, 2025, "
    "doi: 10.1007/s44443-025-00182-4.",
]
INDENT = Cm(1.27)          # text indent of the template's numbered items
CONTENT_W = Cm(14.6)       # page width minus margins minus INDENT


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def add_bullet_numbering(doc):
    """Add a real bullet list definition to numbering.xml; return its numId."""
    numbering = doc.part.numbering_part.element
    abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abs_id, num_id = max(abs_ids + [0]) + 1, max(num_ids + [0]) + 1
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">'
        '<w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        '<w:lvlText w:val=""/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="1080" w:hanging="360"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr>'
        '</w:lvl></w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)      # abstractNums must precede nums
    else:
        numbering.append(abstract)
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abs_id}"/></w:num>'))
    return num_id


class Writer:
    """Builds paragraphs at the end of the body and moves them after an anchor."""

    def __init__(self, doc, anchor, bullet_num):
        self.doc, self.anchor, self.bullet_num = doc, anchor._p, bullet_num

    def _place(self, element):
        self.anchor.addnext(element)
        self.anchor = element

    def para(self, runs, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, after=6,
             keep_next=False, size=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        pf.space_after = Pt(after)
        pf.keep_with_next = keep_next
        if indent:
            pf.left_indent = INDENT
        if isinstance(runs, str):
            runs = [(runs, "")]
        for text, fmt in runs:
            r = p.add_run(text)
            r.bold = "b" in fmt
            r.italic = "i" in fmt
            r.font.subscript = "s" in fmt
            if size:
                r.font.size = Pt(size)
        self._place(p._p)
        return p

    def bullet(self, runs):
        p = self.para(runs, after=2)
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{self.bullet_num}"/></w:numPr>'))
        p.paragraph_format.left_indent = None   # the list level sets the indent
        return p

    def figure(self, png, width, caption_no, caption):
        p = self.para("", align=WD_ALIGN_PARAGRAPH.CENTER, after=2, keep_next=True)
        p.add_run().add_picture(str(png), width=width)
        self.para([(f"Gambar {caption_no}. ", "b"), (caption, "")],
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=10, size=10)

    def table(self, header, rows, widths, caption_no, caption):
        self.para([(f"Tabel {caption_no}. ", "b"), (caption, "")],
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=4, keep_next=True, size=10)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblPr = t._tbl.tblPr
        tblPr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{int(INDENT.emu // 635)}" w:type="dxa"/>'))
        for r_i, values in enumerate([header] + rows):
            trPr = t.rows[r_i]._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if r_i == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for c_i, val in enumerate(values):
                cell = t.cell(r_i, c_i)
                cell.width = widths[c_i]
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                if c_i == 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = par.add_run(val)
                run.font.size = Pt(9.5)
                if r_i == 0:
                    run.bold = True
                    cell._tc.get_or_add_tcPr().append(parse_xml(
                        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9E2F3"/>'))
                elif c_i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        self._place(t._tbl)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        self._place(spacer._p)


def set_cell(cell, lines, highlight=False, bold=False):
    paras = cell.paragraphs
    for i, line in enumerate(lines):
        p = paras[i] if i < len(paras) else cell.add_paragraph()
        r = p.add_run(line)
        r.bold = bold
        if highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


# ---------------------------------------------------------------------------
# content
# ---------------------------------------------------------------------------

def build(template):
    doc = Document(template)
    bullet_num = add_bullet_numbering(doc)

    # identity table
    t = doc.tables[0]
    set_cell(t.cell(0, 1), ["[Nama lengkap] / [NIM]"], highlight=True)
    set_cell(t.cell(1, 1), [JUDUL], bold=True)
    set_cell(t.cell(2, 1), ["1. [Nama Pembimbing 1]", "2. [Nama Pembimbing 2]"], highlight=True)

    paras = doc.paragraphs
    item1 = next(p for p in paras if p.text.startswith("Gambaran system yang sudah"))
    item2 = next(p for p in paras if p.text.startswith("Gambaran system yang diusulkan"))

    # drop the empty spacer paragraphs the template left under both items
    body = list(doc.element.body)
    i1, i2 = body.index(item1._p), body.index(item2._p)
    for el in body[i1 + 1:i2] + body[i2 + 1:]:
        if el.tag == qn("w:p") and not "".join(el.itertext()).strip():
            el.getparent().remove(el)

    # ---------------- 1. sistem yang sudah / pernah ada ----------------
    w = Writer(doc, item1, bullet_num)
    w.para("Smart market yang menjadi objek TA ini belum memiliki sistem penentuan posisi apa "
           "pun; sistem pada TA ini adalah yang pertama dibangun di sana. Troli belanja masih "
           "berupa troli biasa yang didorong pelanggan tanpa perangkat elektronik. Pelanggan "
           "mencari letak produk sendiri melalui papan petunjuk, berkeliling lorong, atau "
           "bertanya kepada petugas, dan pengelola toko tidak dapat mengetahui posisi troli "
           "kecuali dengan melihatnya langsung (Gambar 1).")
    w.para("Akibatnya, dua layanan yang dirancang tim Smart Market belum dapat diwujudkan, "
           "yaitu navigasi menuju rak pada aplikasi dan Smart Trolley yang bergerak secara "
           "otonom. Keduanya membutuhkan koordinat troli secara real-time dengan galat yang "
           "cukup kecil untuk membedakan lorong antar-rak yang lebarnya sekitar 1–2 m.")
    w.para("GPS tidak dapat dipakai karena sinyal satelit teredam oleh bangunan. Penentuan "
           "posisi berbasis kekuatan sinyal WiFi atau Bluetooth Low Energy (RSSI) umumnya hanya "
           "mencapai akurasi orde meter, dan sistem berbasis sensor inframerah pada troli "
           "supermarket melaporkan galat rata-rata hingga sekitar 1 m saat troli bergerak [1].")
    w.para([("Ultra-Wideband (UWB) mengukur jarak dari waktu tempuh sinyal (", ""),
            ("two-way ranging", "i"),
            (") sehingga dapat mencapai akurasi orde sentimeter, dan sudah pernah diteliti "
             "untuk troli belanja. Vandermeeren dan Steendam [2] memasang tag UWB dan sensor "
             "IMU pada troli; galat rata-rata 62,6 cm dengan UWB saja turun menjadi 34,1 cm "
             "dengan Kalman filter. Duan dkk. [3] memakai UWB untuk melacak troli dan "
             "menampilkan panduan belanja di ponsel pelanggan, sedangkan Liang dkk. [4] "
             "memadukan UWB dan kamera dalam SLAM agar troli dapat mengikuti pelanggan di "
             "supermarket tanpa kasir. Penelitian tersebut menunjukkan UWB layak dipakai pada "
             "troli belanja, tetapi masing-masing dibangun untuk sistem dan perangkatnya "
             "sendiri, sehingga smart market ini tetap memerlukan subsistem positioning yang "
             "dirancang untuk Smart Trolley dan aplikasi milik tim.", "")])
    w.para("Kotak merah bernomor pada Gambar 1 menandai bagian yang berubah pada sistem yang "
           "diusulkan:", keep_next=True, after=2)
    w.bullet([("① Area toko: ", "b"),
              ("belum ada perangkat penentu posisi, dan GPS tidak menjangkau dalam ruangan.", "")])
    w.bullet([("② Troli: ", "b"),
              ("troli biasa tanpa perangkat yang dapat mengukur jarak atau mengirim data.", "")])
    w.bullet([("③ Posisi troli: ", "b"),
              ("tidak ada yang menghitung koordinat troli.", "")])
    w.bullet([("④ Pencarian produk: ", "b"),
              ("dilakukan sendiri oleh pelanggan; navigasi pada aplikasi dan troli otonom tidak "
               "dapat dibuat tanpa koordinat troli.", "")])
    w.bullet([("⑤ Pengelolaan: ", "b"),
              ("pengelola tidak memiliki peta maupun pemantauan posisi troli.", "")])
    w.para("", after=2)
    w.figure(HERE / "1_sistem_saat_ini.png", CONTENT_W, 1,
             "Gambaran kondisi saat ini, sebelum ada sistem penentuan posisi. Kotak merah "
             "putus-putus bernomor menandai bagian yang diubah pada sistem yang diusulkan.")

    # ---------------- 2. sistem yang diusulkan ----------------
    w = Writer(doc, item2, bullet_num)
    w.para("Sistem yang diusulkan adalah subsistem indoor positioning berbasis UWB yang "
           "menjawab satu pertanyaan bagi Smart Trolley: di mana troli berada saat ini. "
           "Empat anchor UWB berbiaya rendah (ESP32 + DW1000) dipasang pada posisi tetap yang "
           "diukur di "
           "area smart market, dan satu tag UWB dipasang pada troli. Tag mengukur jarak ke "
           "setiap anchor dengan two-way ranging, lalu mengirim setiap sampel jarak ke server "
           "positioning melalui WiFi. Server mengestimasi posisi troli menggunakan Extended "
           "Kalman Filter (EKF) dan menyediakan posisi tersebut bagi sistem navigasi troli dan "
           "aplikasi Smart Market yang dikembangkan anggota tim lain (Gambar 2).")
    w.para("Bagian yang diubah ditandai dengan kotak merah bernomor yang sama dengan Gambar 1, "
           "dan dirangkum pada Tabel 1.", after=8)
    w.figure(HERE / "2_sistem_usulan.png", CONTENT_W, 2,
             "Gambaran sistem yang diusulkan. Kotak biru adalah ruang lingkup TA ini; kotak "
             "abu-abu adalah TA anggota tim lain yang menerima keluaran posisi.")

    widths = [Cm(1.0), Cm(2.9), Cm(4.4), Cm(6.3)]
    w.table(["No", "Bagian", "Kondisi saat ini", "Sistem yang diusulkan"], [
        ["①", "Area toko",
         "Belum ada perangkat penentu posisi; GPS tidak menjangkau dalam ruangan",
         "Empat anchor UWB (ESP32 + DW1000) dipasang pada posisi tetap yang diukur dengan "
         "meteran laser"],
        ["②", "Troli",
         "Troli biasa tanpa perangkat elektronik",
         "Tag UWB pada troli mengukur jarak ke setiap anchor (two-way ranging); firmware "
         "mengirim setiap sampel jarak mentah dengan timestamp (±5 Hz) melalui WiFi dan dapat "
         "diperbarui lewat OTA"],
        ["③", "Posisi troli",
         "Tidak diketahui; tidak ada yang menghitung koordinat troli",
         "Server positioning: EKF dengan state [x, y, vx, vy, b] dan model jarak miring 3D, "
         "koreksi bias jarak (kalibrasi per anchor + bias online), penyaring outlier dan NLOS "
         "(gating 6σ, bobot Huber)"],
        ["④", "Pencarian produk dan navigasi",
         "Pelanggan mencari produk sendiri; navigasi dan troli otonom belum dapat dibuat",
         "API posisi standar (REST /api/v1/navigation/position dan MQTT 10 Hz) berisi x, y, "
         "quality, valid, serta denah dan obstacle untuk navigasi Smart Trolley dan aplikasi "
         "Smart Market"],
        ["⑤", "Pengelolaan",
         "Tidak ada peta maupun pemantauan troli",
         "Web dashboard: peta 3D posisi troli, pengaturan denah dan anchor, kalibrasi, tuning, "
         "OTA, perekaman sesi dan evaluasi akurasi (CEP, RMSE, cross-track error)"],
    ], widths, 1, "Letak perubahan antara kondisi saat ini dan sistem yang diusulkan")

    w.para([("Ruang lingkup. ", "b"),
            ("TA ini mencakup subsistem positioning: firmware anchor dan tag, server "
             "positioning dengan EKF, antarmuka posisi, serta web dashboard. Navigasi otonom "
             "(path planning, kendali motor, penghindaran halangan) dan aplikasi pengguna "
             "merupakan TA anggota tim lain dan hanya menerima keluaran posisi.", "")])
    w.para([("Parameter yang diuji:", "b")], keep_next=True, after=2)
    w.bullet("galat posisi statis (mean error, RMSE, CEP50, CEP95) pada titik acuan yang "
             "diukur dengan meteran laser;")
    w.bullet("galat lintasan (cross-track error) saat troli bergerak pada lintasan yang "
             "diketahui;")
    w.bullet("laju pembaruan dan latensi posisi;")
    w.bullet("pengaruh NLOS (halangan tubuh manusia dan obstacle) terhadap galat posisi;")
    w.bullet("kestabilan posisi saat troli diam maupun bergerak.")
    w.para("", after=2)

    w.para("Gambar 3 menunjukkan alur pengolahan di server. Setiap sampel jarak hanya dipakai "
           "satu kali dan diproses pada waktu pengukurannya. Sampel yang tidak wajar dibuang, "
           "jaraknya dikoreksi dengan hasil kalibrasi, lalu EKF melakukan prediksi ke waktu "
           "sampel dan update menggunakan model jarak miring", keep_next=True, after=4)
    w.para([("h(x) = √((x − x", "i"), ("a", "is"), (")² + (y − y", "i"), ("a", "is"),
            (")² + Δz²) + b", "i")],
           align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True, after=4)
    w.para([("dengan (x", ""), ("a", "s"), (", y", ""), ("a", "s"),
            (") posisi anchor, Δz beda tinggi anchor dan tag, dan b bias jarak bersama. "
             "Sampel dengan residu terlalu besar ditolak; bila semua sampel ditolak lebih dari "
             "1,5 detik, posisi diinisialisasi ulang dengan multilaterasi.", "")], after=8)
    w.para("Pendekatan ini dipilih karena pada simulasi lintasan lurus dengan galat radio "
           "yang sama, trilaterasi langsung dari jarak terakhir menghasilkan galat lintasan "
           "(cross-track RMSE) 52,7 cm, sedangkan pengolahan dengan EKF ini 11–13 cm.", after=8)
    w.figure(HERE / "3_alur_ekf.png", Cm(13.2), 3,
             "Alur estimasi posisi per sampel jarak pada server positioning.")

    w.para([("Referensi", "b")], keep_next=True, after=4)
    for ref in REFERENSI:
        p = w.para(ref, align=WD_ALIGN_PARAGRAPH.LEFT, after=2, size=9.5)
        p.paragraph_format.left_indent = INDENT + Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)


    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print("wrote", build(sys.argv[1]))
